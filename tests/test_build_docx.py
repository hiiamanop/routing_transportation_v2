import unittest
from pathlib import Path

from scripts.build_docx_manuscripts import convert_markdown_to_docx

ROOT = Path(__file__).resolve().parent.parent
ID_MD = ROOT / "docs" / "manuscript" / "id" / "manuscript.md"
EN_MD = ROOT / "docs" / "manuscript" / "en" / "manuscript.md"
ID_DOCX = ROOT / "docs" / "manuscript" / "id" / "manuscript.docx"
EN_DOCX = ROOT / "docs" / "manuscript" / "en" / "manuscript.docx"


class DocxConversionTests(unittest.TestCase):
    def test_converts_id_manuscript_to_docx_with_tables_and_figures(self):
        import docx
        out_path = convert_markdown_to_docx(ID_MD, ID_DOCX)
        self.assertTrue(out_path.exists())
        self.assertGreater(out_path.stat().st_size, 50000)  # Contains embedded images
        doc = docx.Document(out_path)
        self.assertGreaterEqual(len(doc.paragraphs), 40)
        self.assertEqual(len(doc.tables), 5)

        # Verify inline images exist in document package
        rels = doc.part.rels
        image_parts = [rel.target_ref for rel in rels.values() if "image" in rel.target_ref]
        self.assertEqual(len(image_parts), 6)

    def test_converts_en_manuscript_to_docx_with_tables_and_figures(self):
        import docx
        out_path = convert_markdown_to_docx(EN_MD, EN_DOCX)
        self.assertTrue(out_path.exists())
        self.assertGreater(out_path.stat().st_size, 50000)
        doc = docx.Document(out_path)
        self.assertGreaterEqual(len(doc.paragraphs), 40)
        self.assertEqual(len(doc.tables), 5)

        rels = doc.part.rels
        image_parts = [rel.target_ref for rel in rels.values() if "image" in rel.target_ref]
        self.assertEqual(len(image_parts), 6)


if __name__ == "__main__":
    unittest.main()
